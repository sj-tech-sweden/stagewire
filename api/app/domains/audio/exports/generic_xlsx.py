import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .base import ConsoleExporter, ExportResult
from .common import (
    INPUT_HEADER,
    OUTPUT_HEADER,
    channel_row,
    output_row,
    safe_filename,
)

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")


def _build_sheet(ws, channel_list, header, rows):
    ws.cell(row=1, column=1, value="Stagewire channel list export").font = TITLE_FONT
    ws.cell(row=2, column=1, value=f"Name: {channel_list.name}")
    ws.cell(
        row=3,
        column=1,
        value=f"Console: {channel_list.mixer_product_name or 'unspecified'}",
    )
    header_row_idx = 5
    for c, h in enumerate(header, start=1):
        ws.cell(row=header_row_idx, column=c, value=h)
    data_start = header_row_idx + 1
    for i, r in enumerate(rows):
        for c, v in enumerate(r, start=1):
            ws.cell(row=data_start + i, column=c, value=v)
    data_end = data_start + len(rows) - 1
    ncols = len(header)

    for c in range(1, ncols + 1):
        values = [str(header[c - 1])]
        for r in range(data_start, data_end + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None and v != "":
                values.append(str(v))
        width = min(max(len(x) for x in values) + 2, 48)
        ws.column_dimensions[get_column_letter(c)].width = width

    if data_end >= data_start:
        last_col = get_column_letter(ncols)
        ref = f"A{header_row_idx}:{last_col}{data_end}"
        table = Table(displayName=ws.title.replace(" ", "_"), ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)
        ws.freeze_panes = ws.cell(row=data_start, column=1)
    else:
        for c in range(1, ncols + 1):
            cell = ws.cell(row=header_row_idx, column=c)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL

    for c in range(1, ncols + 1):
        ws.cell(row=header_row_idx, column=c).alignment = Alignment(vertical="center")

    ws.sheet_view.showGridLines = False


class GenericXlsxExporter(ConsoleExporter):
    key = "generic_xlsx"
    label = "Generic Excel (XLSX)"
    description = (
        "Console-agnostic Excel workbook with a filterable, sortable table of "
        "inputs and a separate table of outputs."
    )
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    file_extension = ".xlsx"

    def export(self, channel_list, channels, outputs) -> ExportResult:
        wb = Workbook()
        ws = wb.active
        ws.title = "Inputs"
        _build_sheet(
            ws,
            channel_list,
            INPUT_HEADER,
            [channel_row(c) for c in sorted(channels, key=lambda x: x.channel_number or 0)],
        )
        ws2 = wb.create_sheet("Outputs")
        _build_sheet(
            ws2,
            channel_list,
            OUTPUT_HEADER,
            [output_row(o) for o in sorted(outputs, key=lambda x: x.position or 0)],
        )
        buf = io.BytesIO()
        wb.save(buf)
        return ExportResult(
            content=buf.getvalue(),
            filename=f"{safe_filename(channel_list.name)}.xlsx",
            media_type=self.media_type,
        )
